#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SOAT VAI SO LE — doc so trong so_can_check.txt, do trong Sheet, gui Telegram

Dung khi can kiem tra mot vai so cu the (khach vua nhan trong nhom, hoac
anh Son muon test), thay vi quet toan bo Sheet.

Cach dung: ghi so vao so_can_check.txt (moi dong mot so), day len GitHub.
Workflow tu chay va gui bao cao ve Telegram theo dung mau bot cu.

Chuong trinh CHI DOC Sheet, khong sua gi.
"""

import os
import sys
from datetime import datetime, timedelta, timezone

import requests
from openpyxl import load_workbook

from phone_utils import extract_phones, normalize

SHEET_ID = os.environ.get(
    "SHEET_ID", "1CaCQZvLhAafJ42vuinOlwGhlGmPZCCSzPyzxmlIFNyE")
TG_TOKEN = os.environ.get("TG_TOKEN", "")
TG_CHAT = os.environ.get("TG_CHAT", "")

THU_MUC = os.path.dirname(os.path.abspath(__file__))
FILE_SO = os.path.join(THU_MUC, "so_can_check.txt")
FILE_TAM = "/tmp/don_hang.xlsx"
VN = timezone(timedelta(hours=7))
NGAN = "────────────"

# Ten cot co the gap trong Sheet (chu thuong de so sanh)
COT_SDT = ["sđt", "sdt", "số điện thoại", "số đt", "so dt"]
COT_KHACH = ["khách hàng", "tên khách", "tên", "khách"]
COT_NGAY = ["ngày", "ngày chốt", "ngày gửi"]
COT_SP = ["sản phẩm", "combo"]
COT_TT = ["trạng thái", "tình trạng", "tình trạng khách"]


def gio_hien_tai():
    return datetime.now(VN).strftime("%H:%M %d/%m/%Y")


def doc_so_can_check():
    """Doc danh sach so tu file, bo dong trong va dong ghi chu."""
    if not os.path.exists(FILE_SO):
        return []
    khoa = []
    with open(FILE_SO, encoding="utf-8") as f:
        for dong in f:
            dong = dong.split("#")[0].strip()
            if not dong:
                continue
            for k in extract_phones(dong) or []:
                if k not in khoa:
                    khoa.append(k)
            k = normalize(dong)
            if k and k not in khoa:
                khoa.append(k)
    return khoa


def tai_sheet():
    r = requests.get(
        "https://docs.google.com/spreadsheets/d/%s/export?format=xlsx" % SHEET_ID,
        timeout=90)
    if r.status_code != 200:
        raise SystemExit("Khong tai duoc Sheet (HTTP %d). Sheet da chia se "
                         "'Bat ky ai co link -> Nguoi xem' chua?" % r.status_code)
    with open(FILE_TAM, "wb") as f:
        f.write(r.content)
    return FILE_TAM


def tim_cot(hang, ten_co_the):
    thap = [str(o).strip().lower() if o is not None else "" for o in hang]
    for ten in ten_co_the:
        if ten in thap:
            return thap.index(ten)
    return None


def nap_chi_muc(duong_dan):
    """Tra ve {khoa9so: [ {customer,date,product,status,tab}, ... ]}."""
    wb = load_workbook(duong_dan, read_only=True, data_only=True)
    chi_muc = {}

    for ten_tab in wb.sheetnames:
        ws = wb[ten_tab]
        hang_list = list(ws.iter_rows(values_only=True))

        # tim dong tieu de: dong dau tien co o ten giong cot SDT
        i_hdr = None
        for i, hang in enumerate(hang_list[:30]):
            if tim_cot(hang, COT_SDT) is not None:
                i_hdr = i
                break
        if i_hdr is None:
            continue

        hdr = hang_list[i_hdr]
        i_sdt = tim_cot(hdr, COT_SDT)
        i_khach = tim_cot(hdr, COT_KHACH)
        i_ngay = tim_cot(hdr, COT_NGAY)
        i_sp = tim_cot(hdr, COT_SP)
        i_tt = tim_cot(hdr, COT_TT)

        def o(hang, idx):
            if idx is None or idx >= len(hang) or hang[idx] is None:
                return ""
            return str(hang[idx]).strip()

        for hang in hang_list[i_hdr + 1:]:
            # nhat so o cot SDT, va ca so nam lan trong o khac
            trong_dong = []
            k = normalize(o(hang, i_sdt))
            if k:
                trong_dong.append(k)
            for idx, gia_tri in enumerate(hang):
                if idx == i_sdt or gia_tri is None:
                    continue
                for k2 in extract_phones(str(gia_tri)):
                    if k2 not in trong_dong:
                        trong_dong.append(k2)
            if not trong_dong:
                continue

            ban_ghi = {
                "customer": o(hang, i_khach) or "—",
                "date": o(hang, i_ngay) or "—",
                "product": o(hang, i_sp) or "—",
                "status": o(hang, i_tt) or "—",
                "tab": ten_tab,
            }
            for k3 in trong_dong:
                chi_muc.setdefault(k3, []).append(ban_ghi)

    wb.close()
    return chi_muc


def dung_tin(khoa, cac_dong):
    """Dung tin bao cao dung mau bot DiLi."""
    so = "0" + khoa
    if not cac_dong:
        return ("❌ Không trùng\n"
                "📞 %s\n"
                "(Đã đối chiếu toàn bộ các tab trong Sheet, không thấy số này)\n"
                "🕒 %s" % (so, gio_hien_tai()))

    dong = ["⚠️ TRÙNG ĐƠN — số đã có trong Sheet (%d dòng)" % len(cac_dong),
            "📞 %s" % so, NGAN]
    for i, m in enumerate(cac_dong, 1):
        dong.append("#%d Khách: %s | Ngày: %s | SẢN PHẨM: %s | "
                    "Trạng thái: %s | 📄 Tab: %s"
                    % (i, m["customer"], m["date"], m["product"],
                       m["status"], m["tab"]))
        dong.append(NGAN)
    dong.append("📌 Xử lý trùng đơn DiLi Supplement · 🕒 %s" % gio_hien_tai())
    return "\n".join(dong)


def gui_telegram(chu):
    r = requests.post(
        "https://api.telegram.org/bot%s/sendMessage" % TG_TOKEN,
        data={"chat_id": TG_CHAT, "text": chu,
              "disable_web_page_preview": "true"},
        timeout=30)
    kq = r.json()
    if not kq.get("ok"):
        raise SystemExit("Telegram tu choi: %s" % kq.get("description"))


def main():
    thieu = [t for t, g in (("TG_TOKEN", TG_TOKEN), ("TG_CHAT", TG_CHAT)) if not g]
    if thieu:
        sys.exit("Chua dien ma bi mat: %s" % " ".join(thieu))

    can_check = doc_so_can_check()
    if not can_check:
        print("Khong co so nao trong so_can_check.txt — khong lam gi.")
        return

    print("Can soat %d so: %s" % (len(can_check),
                                  ", ".join("0" + k for k in can_check)))

    chi_muc = nap_chi_muc(tai_sheet())
    print("Sheet co %d so khac nhau." % len(chi_muc))

    trung = 0
    for khoa in can_check:
        cac_dong = chi_muc.get(khoa, [])
        if cac_dong:
            trung += 1
        gui_telegram(dung_tin(khoa, cac_dong))
        print("  0%s -> %s" % (khoa, "TRUNG (%d dong)" % len(cac_dong)
                               if cac_dong else "khong trung"))

    print("XONG. %d/%d so bi trung." % (trung, len(can_check)))


if __name__ == "__main__":
    main()
